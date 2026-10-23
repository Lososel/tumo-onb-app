"""Two batches in one workbook: "4 поток" and "5 поток", red (left) rows, students who moved.

All names and emails are fake.
"""

import io

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from app.services.sheets import _is_red, batch_number, parse_tabs, xlsx_to_tabs

RED = PatternFill("solid", fgColor="FFFF0000")
LIGHT_RED = PatternFill("solid", fgColor="FFF4CCCC")
YELLOW = PatternFill("solid", fgColor="FFFFF2CC")


def workbook() -> bytes:
    wb = openpyxl.Workbook()
    flow4 = wb.active
    flow4.title = "4 поток"
    flow4.append(["ФИО ученика", "Расписание", "Коуч", "Почта TUMO"])
    flow4.append(["Иванова Анна", "Пн/Чт 10:30-12:30", "Alinur", "anna.ivanova@tumo.world"])
    flow4.append(["Петров Олег", "Вт/Пт 14:30-16:30", "Nazym", "oleg.petrov@tumo.world"])  # moved to 5
    flow4.append(["Сидорова Мария", "Пн/Чт 16:30-18:30", "Dana", "maria.sidorova@tumo.world"])  # left
    flow4.append(["Жұмабек Әлия", "Вт/Пт 10:30-12:30", "Aliya", "aliya.zh@tumo.world"])  # highlighted yellow
    for cell in flow4[3]:
        cell.fill = RED
    for cell in flow4[4]:
        cell.fill = LIGHT_RED
    for cell in flow4[5]:
        cell.fill = YELLOW

    flow5 = wb.create_sheet("5 поток")
    # The newer tab splits the schedule into days and time.
    flow5.append(["Name", "Дни", "Время", "TUMO email", "Coach"])
    flow5.append(["Петров Олег", "Пн/Чт", "16:30-18:30", "oleg.petrov@tumo.world", "Samat"])
    flow5.append(["Нұрланов Арман", "Вт/Пт", "10:30-12:30", "arman.n@tumo.world", "Moldir"])
    flow5.append(["Иванова Анна", "Вт/Пт", "18:30-20:30", "", "Gulim"])  # left batch 5 (name cell red)
    flow5["A4"].fill = RED
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def learners():
    snap = parse_tabs(xlsx_to_tabs(workbook()), source="google")
    return {r.full_name: r for r in snap.learners}


def test_both_batches_are_included(learners):
    assert set(learners) == {
        "Иванова Анна", "Петров Олег", "Сидорова Мария", "Жұмабек Әлия", "Нұрланов Арман",
    }
    assert learners["Нұрланов Арман"].tab == "5 поток"
    assert learners["Жұмабек Әлия"].tab == "4 поток"


def test_split_days_and_time_columns_become_the_schedule(learners):
    arman = learners["Нұрланов Арман"]
    assert arman.schedule == "Вт/Пт 10:30-12:30"
    assert arman.coach_name == "Moldir"
    assert arman.default_email == "arman.n@tumo.world"
    assert arman.status_code == "active_schedule"


def test_red_rows_are_inactive_other_fills_are_not(learners):
    assert learners["Сидорова Мария"].status_code == "inactive"  # light red whole row
    assert learners["Жұмабек Әлия"].status_code == "active_schedule"  # yellow is just a highlight


def test_student_in_both_batches_shows_the_newest_batch(learners):
    oleg = learners["Петров Олег"]
    assert oleg.tab == "5 поток"
    assert (oleg.schedule, oleg.coach_name) == ("Пн/Чт 16:30-18:30", "Samat")


def test_active_row_beats_a_red_row_in_a_newer_batch(learners):
    anna = learners["Иванова Анна"]  # red in "5 поток", still active in "4 поток"
    assert anna.active and anna.tab == "4 поток"
    assert (anna.schedule, anna.coach_name) == ("Пн/Чт 10:30-12:30", "Alinur")


def test_lookup_answers_inactive_for_a_red_only_student(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from app.core.config import Settings
    from app.main import create_app
    from app.sources import GoogleSheetsSource

    monkeypatch.setattr(GoogleSheetsSource, "fetch", lambda self: xlsx_to_tabs(workbook()))
    settings = Settings(data_source="google", google_sheet_id="id", google_credentials_json="{}",
                        cache_file=tmp_path / "c.json", lookup_rate_limit_per_minute=0)
    with TestClient(create_app(settings)) as c:
        ask = lambda q: c.post("/api/schedule/lookup", json={"query": q}).json()  # noqa: E731
        assert ask("Сидорова Мария")["status"] == "inactive"
        (card,) = ask("Петров Олег")["results"]
        assert card["schedule"] == "Пн/Чт 16:30-18:30"
        assert ask("Нурланов Арман")["status"] == "ok"  # Kazakh letters folded


@pytest.mark.parametrize("rgb, red", [
    ("FFFF0000", True), ("FFEA9999", True), ("FFF4CCCC", True), ("FFFFC7CE", True), ("FFCC0000", True),
    ("FFFF9900", False), ("FFF9CB9C", False), ("FFFF00FF", False), ("FFFFF2CC", False), ("FFFFFFFF", False),
    ("FF000000", False),
])
def test_red_detection(rgb, red):
    assert _is_red(PatternFill("solid", fgColor=rgb)) is red


def test_batch_number():
    assert [batch_number(t) for t in ("5 поток", "MAIN - 4 Поток (лист ожидания)", "Сводка")] == [5, 4, 0]
