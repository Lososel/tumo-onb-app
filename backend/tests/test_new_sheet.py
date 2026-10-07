"""The live workbook: an .xlsx in Google Drive with "4 поток", "4 поток (лист ожидания)" and more.

All data here is generated and fake.
"""

import json
import sys
import time
import types
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core.config import Settings
from app.main import create_app
from app.services.sheets import SheetsConfigError, match_header, parse_tabs, xlsx_to_tabs
from app.sources import GoogleSheetsSource

HEADER = ["ФИО родителя", "Номер родителя", "почта родителя", "ФИО ребенка", "Номер ребенка",
          "ДР ребенка", "Расписание", "Коуч"]


@pytest.mark.parametrize(
    "header, field",
    [
        ("ФИО ребенка", "full_name"), ("ФИО", "full_name"), ("Full Name", "full_name"),
        ("Расписание", "schedule"), ("Коуч", "coach_name"),
        ("Зона", "room"), ("Learning Zone", "room"), ("Кабинет", "room"),
        ("Почта TUMO", "default_email"), ("Email", "default_email"),
        ("Временный пароль", "temp_password"),
    ],
)
def test_new_sheet_header_variants(header, field):
    assert match_header(header) == field


def test_header_row_below_a_tall_title_block():
    grid = [["TUMO Astana — 4 поток"], [], ["Обновлено: 07.10"], ["Лимит: 60"], [], [], [],
            HEADER, ["Т", "1", "p@example.com", "Сидоров Максим", "2", "03.03.2013", "Вт, Пт 15:00", "Aliya"]]
    (rec,) = parse_tabs({"4 поток": grid}, source="test").learners
    assert (rec.full_name, rec.coach_name) == ("Сидоров Максим", "Aliya")


def make_workbook() -> bytes:
    wb = Workbook()
    main = wb.active
    main.title = "4 поток"
    main.append(HEADER + ["Зона", "Почта TUMO", "Временный пароль"])
    main.append(["Тестова", "8700", "a@example.com", "Сидоров Максим Игоревич", "8701", "03.03.2013",
                 "Вторник, Пятница : 15:00–17:00", "Aliya", "WR 3\n2 этаж", "maxim.s@tumo.example", 1748.0])
    main.append(["Примерова", "8702", "b@example.com", "Жұмабек Әлия", "8703", "04.04.2012",
                 "Понедельник, Четверг : 10:30–12:30", "Nazym", None, None, None])
    wait = wb.create_sheet("4 поток (лист ожидания)")
    wait.append(HEADER)
    wait.append(["Образцова", "8704", "c@example.com", "Кузнецов Даниил", "8705", "05.05.2012",
                 "Понедельник, Четверг : 10:30–12:30", "Dana"])
    wait.append(["Тестова", "8700", "a@example.com", "Сидоров Максим Игоревич", "8701", "03.03.2013",
                 "Суббота", "Dana"])  # also in the main tab
    summary = wb.create_sheet("Сводка")
    summary.append(["Всего учеников", 3])
    other = wb.create_sheet("3 поток")
    other.append(["ФИО ребенка", "Расписание", "Коуч"])
    other.append(["Ахметов Тимур", "Среда 12:00", "Arman"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_all_student_tabs_are_discovered():
    snap = parse_tabs(xlsx_to_tabs(make_workbook()), source="test")
    by_name = {r.full_name: r for r in snap.learners}
    assert set(by_name) == {"Сидоров Максим Игоревич", "Жұмабек Әлия", "Кузнецов Даниил", "Ахметов Тимур"}
    maxim = by_name["Сидоров Максим Игоревич"]
    # Main tab wins over the waitlist tab for a learner listed in both.
    assert (maxim.schedule, maxim.coach_name, maxim.status_code) == (
        "Вторник, Пятница : 15:00–17:00", "Aliya", "active_schedule")
    assert (maxim.room, maxim.default_email, maxim.temp_password) == (
        "WR 3, 2 этаж", "maxim.s@tumo.example", "1748")  # 1748.0 -> "1748"
    assert by_name["Кузнецов Даниил"].status_code == "waitlist"  # from the waitlist tab
    assert by_name["Кузнецов Даниил"].tab == "4 поток (лист ожидания)"


def test_tab_filter_matches_main_and_waitlist_tabs():
    tabs = xlsx_to_tabs(make_workbook(), tab_filter="4 поток")
    assert list(tabs) == ["4 поток", "4 поток (лист ожидания)"]


def test_personal_columns_from_xlsx_are_not_kept():
    dumped = parse_tabs(xlsx_to_tabs(make_workbook()), source="test").model_dump_json()
    for value in ("Тестова", "Примерова", "8700", "8701", "a@example.com", "03.03.2013"):
        assert value not in dumped


def test_garbage_bytes_are_a_config_error():
    with pytest.raises(SheetsConfigError, match="not a readable .xlsx"):
        xlsx_to_tabs(b"definitely not a zip file")


# ---------- Google source: auto-detects Drive .xlsx vs native Sheet ----------

CREDS = json.dumps({"type": "service_account", "client_email": "r@x.iam.gserviceaccount.com",
                    "private_key": "k", "token_uri": "https://oauth2.googleapis.com/token"})


class FakeResponse:
    def __init__(self, payload=None, content=b""):
        self._payload, self.content = payload, content

    def json(self):
        return self._payload


def install_fake_gspread(monkeypatch, mime, workbook=b"", sheet_tabs=None):
    calls = []

    class HTTP:
        def request(self, method, url, params=None, **_):
            calls.append(params)
            if params and params.get("alt") == "media":
                return FakeResponse(content=workbook)
            if mime is None:
                raise RuntimeError("Drive API has not been used in project")
            return FakeResponse({"mimeType": mime})

    class Worksheet:
        def __init__(self, title):
            self.title = title

    class Spreadsheet:
        def worksheets(self):
            return [Worksheet(t) for t in sheet_tabs]

        def values_batch_get(self, ranges):
            return {"valueRanges": [{"values": sheet_tabs[r.strip("'")]} for r in ranges]}

    class Client:
        http_client = HTTP()

        def set_timeout(self, _):
            pass

        def open_by_key(self, key):
            return Spreadsheet()

    monkeypatch.setitem(sys.modules, "gspread",
                        types.SimpleNamespace(service_account_from_dict=lambda info, scopes: Client()))
    return calls


def test_google_source_reads_xlsx_from_drive(monkeypatch):
    calls = install_fake_gspread(monkeypatch, GoogleSheetsSource.XLSX_MIME, workbook=make_workbook())
    tabs = GoogleSheetsSource(None, "1qcmFakeDriveFileId", credentials_json=CREDS).fetch()
    assert "4 поток (лист ожидания)" in tabs
    assert calls[-1]["alt"] == "media"


def test_google_source_reads_native_sheet(monkeypatch):
    install_fake_gspread(monkeypatch, GoogleSheetsSource.SHEET_MIME,
                         sheet_tabs={"4 поток": [["ФИО ребенка", "Коуч"], ["Иван Петров", "Aliya"]]})
    tabs = GoogleSheetsSource(None, "nativeSheetId", credentials_json=CREDS).fetch()
    assert tabs == {"4 поток": [["ФИО ребенка", "Коуч"], ["Иван Петров", "Aliya"]]}


def test_drive_api_unavailable_falls_back_to_sheets_path(monkeypatch):
    install_fake_gspread(monkeypatch, None, sheet_tabs={"T": [["ФИО", "Коуч"], ["Иван Петров", "A"]]})
    assert "T" in GoogleSheetsSource(None, "id", credentials_json=CREDS).fetch()


def test_unsupported_file_type(monkeypatch):
    install_fake_gspread(monkeypatch, "application/pdf")
    with pytest.raises(SheetsConfigError, match="unsupported file type"):
        GoogleSheetsSource(None, "id", credentials_json=CREDS).fetch()


def test_xlsx_end_to_end_with_fallback_fields(monkeypatch, tmp_path):
    install_fake_gspread(monkeypatch, GoogleSheetsSource.XLSX_MIME, workbook=make_workbook())
    settings = Settings(data_source="google", google_sheet_id="1qcmFakeDriveFileId",
                        google_credentials_json=CREDS, cache_file=tmp_path / "c.json")
    with TestClient(create_app(settings)) as c:
        card = c.post("/api/schedule/lookup", json={"query": "алия жумабек"}).json()["results"][0]
        assert card["room"] == card["default_email"] == "Уточняется"
        assert card["temp_password"] is None
        waiting = c.post("/api/schedule/lookup", json={"query": "кузнецов даниил"}).json()["results"][0]
        assert waiting["status_code"] == "waitlist"


# ---------- startup: stale cache cleared, immediate sync ----------


def test_cache_from_another_sheet_is_discarded(tmp_path):
    cache = tmp_path / "c.json"
    old = parse_tabs({"T": [["ФИО", "Коуч"], ["Старый Ученик", "X"]]}, source="google")
    old.origin = "google:OLD_SHEET_ID|"
    cache.write_text(old.model_dump_json(), encoding="utf-8")
    settings = Settings(data_source="google", google_sheet_id="", cache_file=cache)  # new config
    with TestClient(create_app(settings)) as c:
        assert c.get("/api/health").json()["learners_cached"] == 0  # old sheet's data not served
        r = c.post("/api/schedule/lookup", json={"query": "старый ученик"})
        assert r.json()["status"] == "unavailable"
    assert not cache.exists()


def test_startup_syncs_immediately_even_with_a_cache(tmp_path):
    sheet = tmp_path / "sheet.json"
    cache = tmp_path / "c.json"
    sheet.write_text(json.dumps({"T": [["ФИО", "Коуч"], ["Иван Петров", "Aliya"]]}), encoding="utf-8")
    settings = Settings(mock_file=sheet, cache_file=cache)
    with TestClient(create_app(settings)):
        pass  # first run writes the cache
    # The team adds a learner; on restart the new record must appear right away, not after 5 min.
    rows = [["ФИО", "Коуч"], ["Иван Петров", "Aliya"], ["Анна Смирнова", "Dana"]]
    sheet.write_text(json.dumps({"T": rows}), encoding="utf-8")
    with TestClient(create_app(settings)) as c:
        for _ in range(50):
            if c.get("/api/health").json()["learners_cached"] == 2:
                break
            time.sleep(0.05)
        assert c.post("/api/schedule/lookup", json={"query": "анна смирнова"}).json()["status"] == "ok"
